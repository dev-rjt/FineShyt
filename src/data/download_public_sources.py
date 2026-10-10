"""Download and prepare bounded samples of the two public datasets in the brief.

The full London source is about 759 MB compressed and 10 GB uncompressed. This
module uses HTTP range requests to read one partition from the official ZIP,
then stores only an aggregated, non-identifying sample in the repository.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import random
import shutil
import urllib.request
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


UCI_URL = "https://archive.ics.uci.edu/static/public/502/online%2Bretail%2Bii.zip"
LONDON_URL = (
    "https://data.london.gov.uk/download/vqm0d/"
    "04feba67-f1a3-4563-98d0-f3071e3d56d1/Partitioned%20LCL%20Data.zip"
)


class HTTPRangeReader(io.RawIOBase):
    """Seekable HTTP reader for servers that support byte ranges."""

    def __init__(self, url: str):
        self.url = url
        self.headers = {
            "User-Agent": "Mozilla/5.0 (compatible; MegathonDataBuilder/1.0)",
            "Accept": "*/*",
        }
        req = urllib.request.Request(url, method="HEAD", headers=self.headers)
        with urllib.request.urlopen(req, timeout=90) as response:
            self.length = int(response.headers["Content-Length"])
        self.position = 0
        self.cache_start = -1
        self.cache = b""
        self.block_size = 8 * 1024 * 1024

    def seekable(self) -> bool:
        return True

    def readable(self) -> bool:
        return True

    def tell(self) -> int:
        return self.position

    def seek(self, offset: int, whence: int = io.SEEK_SET) -> int:
        if whence == io.SEEK_SET:
            position = offset
        elif whence == io.SEEK_CUR:
            position = self.position + offset
        elif whence == io.SEEK_END:
            position = self.length + offset
        else:
            raise ValueError(f"Unsupported whence: {whence}")
        self.position = max(0, min(position, self.length))
        return self.position

    def read(self, size: int = -1) -> bytes:
        if self.position >= self.length:
            return b""
        wanted = self.length - self.position if size is None or size < 0 else size
        if (
            self.cache_start <= self.position
            and self.position + wanted <= self.cache_start + len(self.cache)
        ):
            offset = self.position - self.cache_start
            payload = self.cache[offset : offset + wanted]
            self.position += len(payload)
            return payload

        fetch_size = max(wanted, self.block_size)
        end = min(self.length - 1, self.position + fetch_size - 1)
        headers = dict(self.headers)
        headers["Range"] = f"bytes={self.position}-{end}"
        request = urllib.request.Request(self.url, headers=headers)
        with urllib.request.urlopen(request, timeout=180) as response:
            self.cache = response.read()
        self.cache_start = self.position
        payload = self.cache[:wanted]
        self.position += len(payload)
        return payload


def download(url: str, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists() and destination.stat().st_size > 0:
        return
    with urllib.request.urlopen(url, timeout=180) as response, destination.open("wb") as out:
        shutil.copyfileobj(response, out)


def prepare_uci(cache_dir: Path, sample_dir: Path, seed: int, rows: int) -> dict:
    archive = cache_dir / "online_retail_ii.zip"
    download(UCI_URL, archive)
    with zipfile.ZipFile(archive) as zf:
        xlsx_names = [n for n in zf.namelist() if n.lower().endswith(".xlsx")]
        if not xlsx_names:
            raise RuntimeError("UCI ZIP contains no XLSX file")
        xlsx_path = cache_dir / "online_retail_II.xlsx"
        if not xlsx_path.exists():
            with zf.open(xlsx_names[0]) as src, xlsx_path.open("wb") as dst:
                shutil.copyfileobj(src, dst)

    rng = random.Random(seed)
    reservoir: list[dict] = []
    seen = 0
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        header = [str(x).strip() if x is not None else "" for x in next(iterator)]
        positions = {name.lower(): idx for idx, name in enumerate(header)}
        for values in iterator:
            customer = values[positions.get("customer id", positions.get("customerid", -1))]
            description = values[positions.get("description", -1)]
            quantity = values[positions.get("quantity", -1)]
            price = values[positions.get("price", positions.get("unitprice", -1))]
            invoice_date = values[positions.get("invoice date", positions.get("invoicedate", -1))]
            country = values[positions.get("country", -1)]
            invoice = values[positions.get("invoice", positions.get("invoiceno", -1))]
            if customer is None or not description or not invoice_date:
                continue
            try:
                amount = float(quantity) * float(price)
            except (TypeError, ValueError):
                continue
            if amount <= 0 or amount > 5000 or str(invoice).upper().startswith("C"):
                continue
            record = {
                "source_customer_key": hashlib.sha256(str(customer).encode()).hexdigest()[:16],
                "invoice_date": invoice_date.isoformat() if hasattr(invoice_date, "isoformat") else str(invoice_date),
                "description": str(description).strip()[:120],
                "quantity": int(quantity),
                "amount_gbp": round(amount, 2),
                "country": str(country),
            }
            seen += 1
            if len(reservoir) < rows:
                reservoir.append(record)
            else:
                candidate = rng.randrange(seen)
                if candidate < rows:
                    reservoir[candidate] = record
    workbook.close()

    output = sample_dir / "uci_online_retail_sample.csv"
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(reservoir[0]))
        writer.writeheader()
        writer.writerows(reservoir)
    return {
        "url": UCI_URL,
        "download_bytes": archive.stat().st_size,
        "valid_rows_seen": seen,
        "sample_rows": len(reservoir),
        "sample_file": str(output),
    }


def _first_csv_member(zf: zipfile.ZipFile) -> str:
    members = [
        info.filename
        for info in zf.infolist()
        if info.filename.lower().endswith(".csv") and info.file_size > 100_000
    ]
    if not members:
        raise RuntimeError("London ZIP contains no usable CSV partition")
    return sorted(members)[0]


def prepare_london(sample_dir: Path, rows: int) -> dict:
    reader = HTTPRangeReader(LONDON_URL)
    with zipfile.ZipFile(reader) as zf:
        member = _first_csv_member(zf)
        with zf.open(member) as binary:
            text = io.TextIOWrapper(binary, encoding="utf-8-sig", errors="replace")
            csv_reader = csv.DictReader(text)
            aggregated: dict[tuple[str, str], list[float]] = defaultdict(list)
            raw_rows = 0
            for row in csv_reader:
                normalized = {str(k).strip().lower(): v for k, v in row.items()}
                household = (
                    normalized.get("lclid")
                    or normalized.get("household_id")
                    or normalized.get("id")
                )
                timestamp = normalized.get("datetime") or normalized.get("date")
                energy = (
                    normalized.get("kwh/hh (per half hour) ")
                    or normalized.get("kwh/hh (per half hour)")
                    or normalized.get("energy(kwh/hh)")
                    or normalized.get("energy")
                )
                if not household or not timestamp or energy in (None, "", "Null"):
                    continue
                try:
                    energy_value = float(energy)
                except ValueError:
                    continue
                month = str(timestamp)[:7]
                aggregated[(str(household), month)].append(energy_value)
                raw_rows += 1
                if raw_rows >= rows:
                    break

    records = []
    for (household, month), readings in sorted(aggregated.items()):
        if not readings:
            continue
        scale_to_month = (48 * 30.4375) / len(readings)
        records.append({
            "source_household_key": hashlib.sha256(household.encode()).hexdigest()[:16],
            "source_month": month,
            "observed_half_hour_rows": len(readings),
            "observed_kwh": round(sum(readings), 4),
            "scaled_monthly_kwh": round(sum(readings) * scale_to_month, 2),
            "mean_half_hour_kwh": round(sum(readings) / len(readings), 6),
        })

    output = sample_dir / "london_smart_meter_monthly_sample.csv"
    with output.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    return {
        "url": LONDON_URL,
        "remote_zip_bytes": reader.length,
        "partition_member": member,
        "raw_rows_read": raw_rows,
        "aggregated_rows": len(records),
        "sample_file": str(output),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cache-dir", type=Path, default=Path("data/external/cache"))
    parser.add_argument("--sample-dir", type=Path, default=Path("data/external/samples"))
    parser.add_argument("--uci-rows", type=int, default=20_000)
    parser.add_argument("--london-rows", type=int, default=250_000)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--only", choices=["all", "uci", "london"], default="all")
    args = parser.parse_args()
    args.cache_dir.mkdir(parents=True, exist_ok=True)
    args.sample_dir.mkdir(parents=True, exist_ok=True)
    metadata_path = args.sample_dir / "source_metadata.json"
    metadata = json.loads(metadata_path.read_text()) if metadata_path.exists() else {}
    if args.only in ("all", "uci"):
        metadata["uci_online_retail_ii"] = prepare_uci(
            args.cache_dir, args.sample_dir, args.seed, args.uci_rows
        )
    if args.only in ("all", "london"):
        metadata["london_smart_meter"] = prepare_london(
            args.sample_dir, args.london_rows
        )
    metadata_path.write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    print(json.dumps(metadata, indent=2))


if __name__ == "__main__":
    main()
